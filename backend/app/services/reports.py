"""Generación de reportes operacionales, de gestión y de auditoría.

Formatos soportados: PDF (ReportLab), DOCX (python-docx) y XLSX (openpyxl).
Los archivos se guardan en storage/reports y se registran en la tabla reports.
"""
from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.security import ensure_aware
from app.models import Alert, AuditLog, Incident, Report, TriageResult, User
from app.models.enums import ReportFormat, ReportType, Severity
from app.schemas import ReportCreate
from app.services.metrics import calculate_overview

SEVERITY_LABELS = {
    "critical": "Crítica",
    "high": "Alta",
    "medium": "Media",
    "low": "Baja",
    "info": "Informativa",
}
STATUS_LABELS = {
    "pending": "Pendiente",
    "triaging": "En triaje",
    "triaged": "Triada",
    "escalated": "Escalada",
    "false_positive": "Falso positivo",
    "closed": "Cerrada",
}
VERDICT_LABELS = {
    "true_positive": "Verdadero positivo",
    "false_positive": "Falso positivo",
    "suspicious": "Sospechoso",
    "benign": "Benigno",
    "": "Sin triaje",
}
REPORT_LABELS = {
    ReportType.OPERATIONAL: "Reporte operacional de alertas",
    ReportType.MANAGEMENT: "Reporte de gestión SOC",
    ReportType.AUDIT: "Reporte de auditoría y trazabilidad",
}


def _naive(value: datetime | None) -> datetime | None:
    aware = ensure_aware(value)
    if aware is None:
        return None
    return aware.astimezone(timezone.utc).replace(tzinfo=None)


def _fmt_dt(value: datetime | None) -> str:
    aware = ensure_aware(value)
    return aware.strftime("%Y-%m-%d %H:%M") if aware else "—"


def _fmt_date(value: datetime | None) -> str:
    aware = ensure_aware(value)
    return aware.strftime("%d/%m/%Y") if aware else "—"


@dataclass
class ReportContext:
    report_type: ReportType
    fmt: ReportFormat
    start: datetime
    end: datetime
    alerts: list[Alert]
    incidents: list[Incident]
    triages: list[TriageResult]
    audit_logs: list[AuditLog]
    generated_by: str
    generated_at: datetime
    overview: Any = None

    @property
    def title(self) -> str:
        return REPORT_LABELS[self.report_type]

    @property
    def period_label(self) -> str:
        return f"{_fmt_date(self.start)} — {_fmt_date(self.end)}"


def _overview(ctx: ReportContext):
    if ctx.overview is None:
        raise ValueError("El contexto del reporte no incluye indicadores calculados")
    return ctx.overview


def _recommendations(overview) -> list[str]:
    kpis = overview.kpis
    result: list[str] = []
    if kpis.critical_open:
        result.append(
            f"Hay {kpis.critical_open} alertas críticas abiertas: priorizar su contención y asignación nominal."
        )
    if kpis.false_positive_rate > 40:
        result.append(
            f"La tasa de falsos positivos es {kpis.false_positive_rate}%: revisar reglas y umbrales de Wazuh para reducir fatiga."
        )
    if kpis.triage_coverage < 80:
        result.append(
            f"La cobertura de triaje es {kpis.triage_coverage}%: revisar el webhook y la cola de reintentos de n8n."
        )
    if kpis.mttt_minutes > 30:
        result.append(
            f"El MTTT es {kpis.mttt_minutes} min: escalar el modelo local o ampliar workers del microservicio."
        )
    if kpis.sla_compliance < 90:
        result.append(
            f"El cumplimiento de SLA es {kpis.sla_compliance}%: priorizar incidentes vencidos y ajustar turnos."
        )
    if not result:
        result.append("Los indicadores se encuentran dentro de los umbrales definidos; mantener el monitoreo.")
    result.append("Conservar la bitácora de auditoría por el periodo de retención definido en la política.")
    return result


def _executive_summary(ctx: ReportContext, overview) -> str:
    k = overview.kpis
    return (
        f"En el periodo {ctx.period_label} se registraron {k.total_alerts} alertas, de las cuales "
        f"{k.open_alerts} permanecen abiertas y {k.critical_open} son críticas. "
        f"La cobertura de triaje automático alcanzó {k.triage_coverage}% y el MTTT fue de "
        f"{k.mttt_minutes} minutos. La tasa de falsos positivos estimada es {k.false_positive_rate}% "
        f"y el cumplimiento de SLA de incidentes es {k.sla_compliance}%. "
        f"Se gestionaron {k.total_incidents} incidentes, con {k.open_incidents} abiertos."
    )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_styles():
    base = getSampleStyleSheet()
    base.add(
        ParagraphStyle(
            name="CoverTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=26,
            leading=30,
            textColor=colors.HexColor("#0f172a"),
            alignment=TA_LEFT,
            spaceAfter=10,
        )
    )
    base.add(
        ParagraphStyle(
            name="CoverBrand",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            textColor=colors.HexColor("#0ea5e9"),
            spaceAfter=24,
        )
    )
    base.add(
        ParagraphStyle(
            name="SectionTitle",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#0f172a"),
            spaceBefore=14,
            spaceAfter=8,
        )
    )
    base.add(
        ParagraphStyle(
            name="BodySmall",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            alignment=TA_JUSTIFY,
            textColor=colors.HexColor("#334155"),
        )
    )
    base.add(
        ParagraphStyle(
            name="Cell",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor("#1e293b"),
        )
    )
    base.add(
        ParagraphStyle(
            name="CellHeader",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.white,
        )
    )
    base.add(
        ParagraphStyle(
            name="KpiLabel",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            textColor=colors.HexColor("#64748b"),
        )
    )
    base.add(
        ParagraphStyle(
            name="KpiValue",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
        )
    )
    return base


def _pdf_header_footer(canvas, doc) -> None:
    canvas.saveState()
    width, height = A4
    canvas.setStrokeColor(colors.HexColor("#e2e8f0"))
    canvas.setLineWidth(0.5)
    canvas.line(1.6 * cm, height - 1.35 * cm, width - 1.6 * cm, height - 1.35 * cm)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.setFillColor(colors.HexColor("#0ea5e9"))
    canvas.drawString(1.6 * cm, height - 1.15 * cm, "SOC-Lite")
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#64748b"))
    canvas.drawRightString(width - 1.6 * cm, height - 1.15 * cm, "Confidencial · Uso interno")
    canvas.line(1.6 * cm, 1.25 * cm, width - 1.6 * cm, 1.25 * cm)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(1.6 * cm, 0.95 * cm, settings.report_brand)
    canvas.drawRightString(width - 1.6 * cm, 0.95 * cm, f"Página {doc.page}")
    canvas.restoreState()


def _kpi_table(overview, styles) -> Table:
    k = overview.kpis
    entries = [
        ("Alertas totales", str(k.total_alerts), "#0ea5e9"),
        ("Alertas abiertas", str(k.open_alerts), "#f59e0b"),
        ("Críticas abiertas", str(k.critical_open), "#ef4444"),
        ("Cobertura de triaje", f"{k.triage_coverage}%", "#10b981"),
        ("MTTT", f"{k.mttt_minutes} min", "#6366f1"),
        ("Falsos positivos", f"{k.false_positive_rate}%", "#8b5cf6"),
        ("MTTR", f"{k.mttr_hours} h", "#0ea5e9"),
        ("Cumplimiento SLA", f"{k.sla_compliance}%", "#10b981"),
        ("Incidentes abiertos", str(k.open_incidents), "#ef4444"),
    ]
    cells = []
    for label, value, color in entries:
        cells.append(
            [
                Paragraph(escape(label), styles["KpiLabel"]),
                Paragraph(f'<font color="{color}">{escape(value)}</font>', styles["KpiValue"]),
            ]
        )
    # 3 columnas x 3 filas de tarjetas
    rows = []
    for index in range(0, len(cells), 3):
        row = []
        for card in cells[index : index + 3]:
            row.append([card[0], card[1]])
        while len(row) < 3:
            row.append(["", ""])
        rows.append(row)
    grid = Table(rows, colWidths=[5.8 * cm] * 3, rowHeights=[1.65 * cm] * 3)
    grid.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e2e8f0")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ]
        )
    )
    return grid


def _severity_chart(overview) -> Drawing:
    drawing = Drawing(460, 170)
    chart = VerticalBarChart()
    chart.x = 34
    chart.y = 28
    chart.width = 410
    chart.height = 120
    chart.data = [[item.value for item in overview.by_severity]]
    chart.categoryAxis.categoryNames = [
        SEVERITY_LABELS.get(item.name, item.name) for item in overview.by_severity
    ]
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 8
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = max(5, max([item.value for item in overview.by_severity] + [1]) + 1)
    chart.valueAxis.valueStep = max(1, (chart.valueAxis.valueMax // 5) or 1)
    chart.bars[0].fillColor = colors.HexColor("#0ea5e9")
    chart.bars[0].strokeColor = None
    chart.barWidth = 12
    drawing.add(chart)
    return drawing


def _trend_chart(overview) -> Drawing:
    from reportlab.graphics.charts.linecharts import HorizontalLineChart

    drawing = Drawing(460, 170)
    chart = HorizontalLineChart()
    chart.x = 34
    chart.y = 30
    chart.width = 410
    chart.height = 115
    values = [point.value for point in overview.alerts_by_day]
    chart.data = [values or [0]]
    step = max(1, (len(overview.alerts_by_day) + 7) // 8)
    labels = []
    for index, point in enumerate(overview.alerts_by_day):
        show = index % step == 0 or index == len(overview.alerts_by_day) - 1
        labels.append(point.label[5:] if show else "")
    chart.categoryAxis.categoryNames = labels
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 6.5
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.boxAnchor = "ne"
    chart.categoryAxis.labels.dx = -2
    chart.categoryAxis.labels.dy = -4
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.lines[0].strokeColor = colors.HexColor("#6366f1")
    chart.lines[0].strokeWidth = 1.6
    drawing.add(chart)
    return drawing


def _table(headers: list[str], rows: list[list[str]], widths: list[float], styles) -> Table:
    data = [[Paragraph(escape(h), styles["CellHeader"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(cell)), styles["Cell"]) for cell in row])
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def build_pdf(ctx: ReportContext) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.9 * cm,
        bottomMargin=1.6 * cm,
        title=ctx.title,
        author="SOC-Lite",
        subject="Reporte de ciberseguridad",
    )
    styles = _pdf_styles()
    overview = _overview(ctx)
    story: list[Any] = []

    # Portada
    story.append(Spacer(1, 3.4 * cm))
    story.append(Paragraph("SOC-Lite · Security Operations Center", styles["CoverBrand"]))
    story.append(Paragraph(ctx.title, styles["CoverTitle"]))
    story.append(Spacer(1, 0.4 * cm))
    story.append(
        Paragraph(
            f"Periodo analizado: <b>{escape(ctx.period_label)}</b>",
            styles["BodySmall"],
        )
    )
    story.append(Paragraph(f"Generado por: <b>{escape(ctx.generated_by)}</b>", styles["BodySmall"]))
    story.append(
        Paragraph(
            f"Fecha de generación: <b>{_fmt_dt(ctx.generated_at)}</b>",
            styles["BodySmall"],
        )
    )
    story.append(Spacer(1, 2.0 * cm))
    story.append(
        Paragraph(
            "Documento confidencial. Contiene datos operacionales de seguridad. "
            "Su distribución está limitada al personal autorizado y a los organismos de control.",
            styles["BodySmall"],
        )
    )
    story.append(PageBreak())

    story.append(Paragraph("1. Resumen ejecutivo", styles["SectionTitle"]))
    story.append(Paragraph(escape(_executive_summary(ctx, overview)), styles["BodySmall"]))
    story.append(Spacer(1, 8))
    story.append(_kpi_table(overview, styles))
    story.append(PageBreak())

    story.append(Paragraph("2. Indicadores y tendencia", styles["SectionTitle"]))
    story.append(Paragraph("Distribución de alertas por severidad", styles["BodySmall"]))
    story.append(_severity_chart(overview))
    story.append(Paragraph("Evolución diaria de alertas", styles["BodySmall"]))
    story.append(_trend_chart(overview))
    story.append(Spacer(1, 6))

    story.append(Paragraph("3. Alertas del periodo", styles["SectionTitle"]))
    rows: list[list[str]] = []
    for alert in ctx.alerts[:30]:
        rows.append(
            [
                _fmt_dt(alert.received_at),
                SEVERITY_LABELS.get(alert.severity.value, alert.severity.value),
                alert.source,
                alert.title,
                STATUS_LABELS.get(alert.status.value, alert.status.value),
                VERDICT_LABELS.get(
                    alert.triage_result.verdict.value if alert.triage_result else "", "Sin triaje"
                ),
            ]
        )
    if not rows:
        rows = [["—", "—", "—", "Sin alertas en el periodo", "—", "—"]]
    story.append(
        _table(
            ["Fecha", "Sev.", "Fuente", "Título", "Estado", "Veredicto"],
            rows,
            [2.55 * cm, 1.35 * cm, 1.65 * cm, 6.7 * cm, 1.9 * cm, 2.35 * cm],
            styles,
        )
    )
    if len(ctx.alerts) > 30:
        story.append(
            Paragraph(
                f"Se muestran 30 de {len(ctx.alerts)} alertas. El detalle completo está disponible en el anexo XLSX.",
                styles["BodySmall"],
            )
        )

    if ctx.report_type in {ReportType.OPERATIONAL, ReportType.MANAGEMENT}:
        story.append(PageBreak())
        story.append(Paragraph("4. Triajes relevantes", styles["SectionTitle"]))
        triage_rows = []
        for alert in ctx.alerts:
            if alert.triage_result is None:
                continue
            triage_rows.append(
                [
                    _fmt_dt(alert.received_at),
                    SEVERITY_LABELS.get(alert.triage_result.recommended_severity.value, ""),
                    VERDICT_LABELS.get(alert.triage_result.verdict.value, ""),
                    f"{alert.triage_result.confidence * 100:.0f}%",
                    alert.triage_result.summary[:260],
                ]
            )
        if not triage_rows:
            triage_rows = [["—", "—", "—", "—", "Sin triajes en el periodo"]]
        story.append(
            _table(
                ["Fecha", "Sev. rec.", "Veredicto", "Conf.", "Resumen del triaje"],
                triage_rows[:15],
                [2.4 * cm, 1.55 * cm, 2.3 * cm, 1.2 * cm, 9.05 * cm],
                styles,
            )
        )

        story.append(Paragraph("5. Incidentes", styles["SectionTitle"]))
        incident_rows = [
            [
                incident.code,
                _fmt_date(incident.opened_at),
                SEVERITY_LABELS.get(incident.severity.value, ""),
                incident.status.value,
                incident.title,
                _fmt_dt(incident.closed_at),
            ]
            for incident in ctx.incidents[:20]
        ]
        if not incident_rows:
            incident_rows = [["—", "—", "—", "—", "Sin incidentes en el periodo", "—"]]
        story.append(
            _table(
                ["Código", "Apertura", "Sev.", "Estado", "Título", "Cierre"],
                incident_rows,
                [2.3 * cm, 1.8 * cm, 1.3 * cm, 2.0 * cm, 6.9 * cm, 2.2 * cm],
                styles,
            )
        )

    story.append(Paragraph("6. Recomendaciones", styles["SectionTitle"]))
    for recommendation in _recommendations(overview):
        story.append(Paragraph(f"• {escape(recommendation)}", styles["BodySmall"]))
        story.append(Spacer(1, 2))

    if ctx.report_type == ReportType.AUDIT:
        story.append(PageBreak())
        story.append(Paragraph("7. Bitácora de auditoría", styles["SectionTitle"]))
        audit_rows = [
            [
                _fmt_dt(log.created_at),
                log.actor_email,
                log.module,
                log.event_type,
                log.action,
                log.direction or "—",
                log.ip_address or "—",
            ]
            for log in ctx.audit_logs[:80]
        ]
        if not audit_rows:
            audit_rows = [["—", "—", "—", "—", "Sin eventos", "—", "—"]]
        story.append(
            _table(
                ["Fecha/hora", "Usuario", "Módulo", "Evento", "Acción", "Dir.", "IP"],
                audit_rows,
                [2.55 * cm, 3.45 * cm, 1.5 * cm, 2.1 * cm, 4.0 * cm, 0.95 * cm, 1.4 * cm],
                styles,
            )
        )

    story.append(Spacer(1, 16))
    story.append(
        Paragraph(
            "Trazabilidad: este reporte se generó desde SOC-Lite y su creación quedó registrada "
            "en la bitácora de auditoría con usuario, fecha, hora y acción.",
            styles["BodySmall"],
        )
    )

    doc.build(story, onFirstPage=_pdf_header_footer, onLaterPages=_pdf_header_footer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def _shade_cell(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _docx_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for index, header in enumerate(headers):
        hdr[index].text = ""
        paragraph = hdr[index].paragraphs[0]
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(255, 255, 255)
        _shade_cell(hdr[index], "0F172A")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(7.5)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width
    return table


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build_docx(ctx: ReportContext) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(9.5)
    for style_name, size, color in (
        ("Title", 24, RGBColor(15, 23, 42)),
        ("Heading 1", 15, RGBColor(15, 23, 42)),
        ("Heading 2", 11, RGBColor(14, 165, 233)),
    ):
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.color.rgb = color

    footer = section.footer.paragraphs[0]
    footer.text = f"{settings.report_brand} · Confidencial · Página "
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(7.5)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    _add_page_number(footer)

    overview = _overview(ctx)
    header = document.add_paragraph()
    run = header.add_run("SOC-LITE · SECURITY OPERATIONS CENTER")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(14, 165, 233)
    document.add_heading(ctx.title, 0)
    document.add_paragraph(f"Periodo: {ctx.period_label}")
    document.add_paragraph(f"Generado por: {ctx.generated_by}")
    document.add_paragraph(f"Fecha de generación: {_fmt_dt(ctx.generated_at)}")
    document.add_paragraph(
        "Documento confidencial. Su distribución está limitada al personal autorizado."
    )

    document.add_heading("1. Resumen ejecutivo", level=1)
    document.add_paragraph(_executive_summary(ctx, overview))

    document.add_heading("2. Indicadores clave", level=1)
    k = overview.kpis
    kpi_rows = [
        ["Alertas totales", str(k.total_alerts), "Alertas abiertas", str(k.open_alerts)],
        ["Críticas abiertas", str(k.critical_open), "Altas abiertas", str(k.high_open)],
        ["Cobertura de triaje", f"{k.triage_coverage}%", "Automatización", f"{k.automation_rate}%"],
        ["MTTT", f"{k.mttt_minutes} min", "MTTR", f"{k.mttr_hours} h"],
        ["Falsos positivos", f"{k.false_positive_rate}%", "SLA", f"{k.sla_compliance}%"],
        ["Incidentes abiertos", str(k.open_incidents), "Incidentes totales", str(k.total_incidents)],
    ]
    _docx_table(document, ["Indicador", "Valor", "Indicador", "Valor"], kpi_rows)

    document.add_heading("3. Distribución por severidad", level=1)
    severity_rows = [
        [SEVERITY_LABELS.get(item.name, item.name), str(item.value)]
        for item in overview.by_severity
    ]
    _docx_table(document, ["Severidad", "Alertas"], severity_rows)

    document.add_heading("4. Alertas del periodo", level=1)
    alert_rows = [
        [
            _fmt_dt(alert.received_at),
            SEVERITY_LABELS.get(alert.severity.value, alert.severity.value),
            alert.source,
            alert.title,
            STATUS_LABELS.get(alert.status.value, alert.status.value),
        ]
        for alert in ctx.alerts[:40]
    ]
    if not alert_rows:
        alert_rows = [["—", "—", "—", "Sin alertas", "—"]]
    _docx_table(document, ["Fecha", "Severidad", "Fuente", "Título", "Estado"], alert_rows)

    if ctx.report_type in {ReportType.OPERATIONAL, ReportType.MANAGEMENT}:
        document.add_heading("5. Triajes y acciones recomendadas", level=1)
        for alert in ctx.alerts[:12]:
            if alert.triage_result is None:
                continue
            triage = alert.triage_result
            document.add_heading(f"{alert.title[:100]}", level=2)
            document.add_paragraph(
                f"Veredicto: {VERDICT_LABELS.get(triage.verdict.value, triage.verdict.value)} · "
                f"Severidad recomendada: {SEVERITY_LABELS.get(triage.recommended_severity.value, '')} · "
                f"Confianza: {triage.confidence * 100:.0f}% · Riesgo: {triage.risk_score}/100"
            )
            document.add_paragraph(triage.summary)
            for action in (triage.recommended_actions or [])[:5]:
                if isinstance(action, dict):
                    document.add_paragraph(
                        f"• [{action.get('priority', '')}] {action.get('action', '')} "
                        f"({action.get('owner', '')}, SLA {action.get('sla_minutes', '')} min)",
                        style="List Bullet",
                    )

        document.add_heading("6. Incidentes", level=1)
        incident_rows = [
            [
                incident.code,
                _fmt_date(incident.opened_at),
                SEVERITY_LABELS.get(incident.severity.value, ""),
                incident.status.value,
                incident.title,
                _fmt_dt(incident.closed_at),
            ]
            for incident in ctx.incidents[:30]
        ]
        if not incident_rows:
            incident_rows = [["—", "—", "—", "—", "Sin incidentes", "—"]]
        _docx_table(
            document,
            ["Código", "Apertura", "Severidad", "Estado", "Título", "Cierre"],
            incident_rows,
        )

    document.add_heading("7. Recomendaciones", level=1)
    for recommendation in _recommendations(overview):
        document.add_paragraph(recommendation, style="List Bullet")

    if ctx.report_type == ReportType.AUDIT:
        document.add_heading("8. Bitácora de auditoría", level=1)
        audit_rows = [
            [
                _fmt_dt(log.created_at),
                log.actor_email,
                log.module,
                log.event_type,
                log.action,
                log.direction or "—",
                log.ip_address or "—",
            ]
            for log in ctx.audit_logs[:120]
        ]
        if not audit_rows:
            audit_rows = [["—", "—", "—", "—", "Sin eventos", "—", "—"]]
        _docx_table(
            document,
            ["Fecha/hora", "Usuario", "Módulo", "Evento", "Acción", "Dir.", "IP"],
            audit_rows,
        )

    document.add_paragraph(
        "Trazabilidad: la generación de este documento quedó registrada en la bitácora de SOC-Lite."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(color="0F172A", bold=True, size=16)
SUBTITLE_FONT = Font(color="0EA5E9", bold=True, size=9)
THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
SEVERITY_FILLS = {
    "critical": PatternFill("solid", fgColor="FEE2E2"),
    "high": PatternFill("solid", fgColor="FFEDD5"),
    "medium": PatternFill("solid", fgColor="FEF3C7"),
    "low": PatternFill("solid", fgColor="DCFCE7"),
    "info": PatternFill("solid", fgColor="E0F2FE"),
}


def _safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_sheet(ws, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_safe(value) for value in row])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def build_xlsx(ctx: ReportContext) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    overview = _overview(ctx)
    k = overview.kpis

    ws["A1"] = "SOC-Lite · Reporte de seguridad"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ctx.title
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = f"Periodo: {ctx.period_label}"
    ws["A4"] = f"Generado por: {ctx.generated_by} · {_fmt_dt(ctx.generated_at)}"
    ws["A5"] = "Documento confidencial de uso interno"

    ws["A7"] = "Indicador"
    ws["B7"] = "Valor"
    for cell in ws[7]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    kpis = [
        ("Alertas totales", k.total_alerts),
        ("Alertas abiertas", k.open_alerts),
        ("Alertas críticas abiertas", k.critical_open),
        ("Alertas altas abiertas", k.high_open),
        ("Tasa de falsos positivos (%)", k.false_positive_rate),
        ("Cobertura de triaje (%)", k.triage_coverage),
        ("Tasa de automatización (%)", k.automation_rate),
        ("MTTT (minutos)", k.mttt_minutes),
        ("MTTR (horas)", k.mttr_hours),
        ("Cumplimiento de SLA (%)", k.sla_compliance),
        ("Incidentes abiertos", k.open_incidents),
        ("Incidentes totales", k.total_incidents),
    ]
    for index, (label, value) in enumerate(kpis, start=8):
        ws.cell(row=index, column=1, value=label)
        ws.cell(row=index, column=2, value=value)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16

    start_row = 22
    ws.cell(row=start_row, column=1, value="Severidad").font = Font(bold=True)
    ws.cell(row=start_row, column=2, value="Alertas").font = Font(bold=True)
    for index, item in enumerate(overview.by_severity, start=start_row + 1):
        ws.cell(row=index, column=1, value=SEVERITY_LABELS.get(item.name, item.name))
        ws.cell(row=index, column=2, value=item.value)

    start_row = start_row + 8
    ws.cell(row=start_row, column=1, value="Fuente").font = Font(bold=True)
    ws.cell(row=start_row, column=2, value="Alertas").font = Font(bold=True)
    for index, item in enumerate(overview.by_source, start=start_row + 1):
        ws.cell(row=index, column=1, value=item.name)
        ws.cell(row=index, column=2, value=item.value)

    # Gráfico de severidad
    pie = PieChart()
    pie.title = "Alertas por severidad"
    data = Reference(ws, min_col=2, min_row=22, max_row=22 + len(overview.by_severity))
    cats = Reference(ws, min_col=1, min_row=23, max_row=22 + len(overview.by_severity))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(cats)
    pie.height = 7
    pie.width = 11
    ws.add_chart(pie, "D7")

    # Hoja de alertas
    ws_alerts = wb.create_sheet("Alertas")
    headers = [
        "ID",
        "Fecha/hora",
        "Fuente",
        "Regla",
        "Nivel",
        "Severidad",
        "Estado",
        "Título",
        "IP origen",
        "Host",
        "Activo",
        "Criticidad activo",
        "Veredicto",
        "Sev. recomendada",
        "Confianza",
        "Riesgo",
        "Resumen triaje",
    ]
    rows = []
    for alert in ctx.alerts:
        triage = alert.triage_result
        rows.append(
            [
                str(alert.id),
                _naive(alert.received_at),
                alert.source,
                alert.rule_id or "",
                alert.rule_level,
                alert.severity.value,
                alert.status.value,
                alert.title,
                alert.source_ip or "",
                alert.hostname or "",
                alert.asset.hostname if alert.asset else "",
                alert.asset.criticality.value if alert.asset else "",
                triage.verdict.value if triage else "",
                triage.recommended_severity.value if triage else "",
                round(triage.confidence * 100, 1) if triage else "",
                triage.risk_score if triage else "",
                triage.summary if triage else "",
            ]
        )
    _write_sheet(
        ws_alerts,
        headers,
        rows,
        [36, 18, 12, 8, 7, 10, 14, 42, 16, 18, 18, 12, 16, 14, 9, 8, 60],
    )
    for row in ws_alerts.iter_rows(min_row=2, max_row=ws_alerts.max_row):
        row[1].number_format = "yyyy-mm-dd hh:mm"
        fill = SEVERITY_FILLS.get(str(row[5].value))
        if fill:
            row[5].fill = fill

    # Hoja de triaje
    ws_triage = wb.create_sheet("Triaje")
    triage_rows = []
    for triage in ctx.triages:
        triage_rows.append(
            [
                str(triage.id),
                str(triage.alert_id),
                triage.verdict.value,
                triage.recommended_severity.value,
                round(triage.confidence * 100, 1),
                triage.risk_score,
                triage.summary,
                triage.reasoning,
                "; ".join(str(item) for item in (triage.mitre_techniques or [])),
                triage.engine,
                triage.model_name,
                triage.latency_ms,
                _naive(triage.created_at),
            ]
        )
    _write_sheet(
        ws_triage,
        [
            "ID",
            "Alerta",
            "Veredicto",
            "Sev. recomendada",
            "Confianza (%)",
            "Riesgo",
            "Resumen",
            "Razonamiento",
            "MITRE",
            "Motor",
            "Modelo",
            "Latencia (ms)",
            "Fecha",
        ],
        triage_rows,
        [36, 36, 16, 15, 12, 9, 60, 60, 24, 16, 20, 12, 18],
    )

    # Hoja de incidentes
    ws_incidents = wb.create_sheet("Incidentes")
    incident_rows = [
        [
            incident.code,
            incident.title,
            incident.severity.value,
            incident.status.value,
            _naive(incident.opened_at),
            _naive(incident.closed_at),
            _naive(incident.sla_due_at),
            incident.owner.full_name if incident.owner else "",
            incident.ticket_external_id or "",
        ]
        for incident in ctx.incidents
    ]
    _write_sheet(
        ws_incidents,
        ["Código", "Título", "Severidad", "Estado", "Apertura", "Cierre", "SLA", "Responsable", "Ticket"],
        incident_rows,
        [16, 50, 12, 16, 18, 18, 18, 24, 18],
    )
    for row in ws_incidents.iter_rows(min_row=2, max_row=ws_incidents.max_row):
        for index in (4, 5, 6):
            row[index].number_format = "yyyy-mm-dd hh:mm"

    # Hoja de auditoría
    ws_audit = wb.create_sheet("Auditoria")
    _write_sheet(
        ws_audit,
        [
            "Fecha/hora",
            "Usuario",
            "Rol",
            "Módulo",
            "Evento",
            "Acción",
            "Dirección",
            "Entidad",
            "IP",
            "Resultado",
        ],
        [
            [
                _naive(log.created_at),
                log.actor_email,
                log.actor_role,
                log.module,
                log.event_type,
                log.action,
                log.direction or "",
                log.entity_type or "",
                log.ip_address or "",
                log.status_code or "",
            ]
            for log in ctx.audit_logs
        ],
        [18, 30, 14, 14, 18, 40, 10, 16, 16, 10],
    )
    for row in ws_audit.iter_rows(min_row=2, max_row=ws_audit.max_row):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"

    # Hoja de conclusiones
    ws_rec = wb.create_sheet("Recomendaciones")
    ws_rec["A1"] = "Recomendaciones y conclusiones"
    ws_rec["A1"].font = TITLE_FONT
    for index, recommendation in enumerate(_recommendations(overview), start=3):
        ws_rec.cell(row=index, column=1, value=f"{index - 2}. {recommendation}")
        ws_rec.cell(row=index, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_rec.column_dimensions["A"].width = 120

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Caso de uso
# ---------------------------------------------------------------------------
def generate_report(db: Session, user: User, request: ReportCreate) -> Report:
    now = datetime.now(timezone.utc)
    if request.period_end:
        end = ensure_aware(request.period_end) or now
    else:
        end = now
    if request.period_start:
        start = ensure_aware(request.period_start) or (end - timedelta(days=request.days))
    else:
        start = end - timedelta(days=request.days)
    if start >= end:
        raise ValueError("El inicio del periodo debe ser anterior al fin")

    alerts = list(
        db.scalars(
            select(Alert)
            .where(Alert.received_at >= start, Alert.received_at <= end)
            .order_by(Alert.received_at.desc())
        ).all()
    )
    incidents = list(
        db.scalars(
            select(Incident)
            .where(Incident.opened_at >= start, Incident.opened_at <= end)
            .order_by(Incident.opened_at.desc())
        ).all()
    )
    triages = list(
        db.scalars(
            select(TriageResult)
            .where(TriageResult.created_at >= start, TriageResult.created_at <= end)
            .order_by(TriageResult.created_at.desc())
        ).all()
    )
    audit_logs: list[AuditLog] = []
    if request.report_type == ReportType.AUDIT:
        audit_logs = list(
            db.scalars(
                select(AuditLog)
                .where(AuditLog.created_at >= start, AuditLog.created_at <= end)
                .order_by(AuditLog.created_at.desc())
                .limit(2000)
            ).all()
        )

    ctx = ReportContext(
        report_type=request.report_type,
        fmt=request.format,
        start=start,
        end=end,
        alerts=alerts,
        incidents=incidents,
        triages=triages,
        audit_logs=audit_logs,
        generated_by=user.full_name,
        generated_at=now,
        overview=calculate_overview(db, start=start, end=end),
    )

    if request.format == ReportFormat.PDF:
        content = build_pdf(ctx)
        extension = "pdf"
    elif request.format == ReportFormat.DOCX:
        content = build_docx(ctx)
        extension = "docx"
    else:
        content = build_xlsx(ctx)
        extension = "xlsx"

    filename = (
        f"SOC-Lite_{request.report_type.value}_{now.strftime('%Y%m%d_%H%M')}_"
        f"{uuid.uuid4().hex[:6]}.{extension}"
    )
    path = settings.reports_path / filename
    path.write_bytes(content)

    report = Report(
        name=f"{REPORT_LABELS[request.report_type]} · {ctx.period_label}",
        report_type=request.report_type,
        format=request.format,
        period_start=start,
        period_end=end,
        filters={"days": request.days},
        file_name=filename,
        size_bytes=len(content),
        generated_by_id=user.id,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
